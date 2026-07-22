from typing import List, Dict, Any, Optional
import docx
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from manuscriptfinesse.export.base_exporter import BaseExporter


class DOCXExporter(BaseExporter):
    """Stage 6 DOCX Exporter using python-docx formatted for standard 12pt manuscript output."""

    def compile(
        self,
        chapters: List[Dict[str, str]],
        output_path: str,
        front_matter: Optional[Dict[str, str]] = None
    ) -> str:
        self._ensure_output_dir(output_path)

        doc = Document()

        # Set standard 1-inch margins
        for section in doc.sections:
            section.top_margin = Inches(1.0)
            section.bottom_margin = Inches(1.0)
            section.left_margin = Inches(1.0)
            section.right_margin = Inches(1.0)

        # Set default font style (Times New Roman 12pt, 1.5 line spacing)
        normal_style = doc.styles['Normal']
        normal_font = normal_style.font
        normal_font.name = 'Times New Roman'
        normal_font.size = Pt(12)
        normal_style.paragraph_format.line_spacing = 1.5
        normal_style.paragraph_format.space_after = Pt(6)

        # Title Page
        title_para = doc.add_paragraph()
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title_para.add_run(self.title)
        title_run.bold = True
        title_run.font.size = Pt(24)

        author_para = doc.add_paragraph()
        author_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        author_run = author_para.add_run(f"By {self.author}")
        author_run.font.size = Pt(14)
        author_run.italic = True

        doc.add_paragraph()  # Blank spacing

        # Front Matter
        if front_matter:
            if "copyright" in front_matter:
                cp_para = doc.add_paragraph(front_matter["copyright"])
                cp_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

            if "dedication" in front_matter:
                doc.add_page_break()
                ded_head = doc.add_heading("Dedication", level=2)
                ded_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
                doc.add_paragraph(front_matter["dedication"])

            if "foreword" in front_matter:
                doc.add_page_break()
                fw_head = doc.add_heading("Foreword", level=2)
                fw_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for block in front_matter["foreword"].split("\n\n"):
                    if block.strip():
                        doc.add_paragraph(block.strip())

        # Chapters
        for i, ch in enumerate(chapters, start=1):
            doc.add_page_break()
            title = ch.get("title", f"Chapter {i}")
            content = ch.get("content", "")

            ch_head = doc.add_heading(title, level=1)
            ch_head.alignment = WD_ALIGN_PARAGRAPH.CENTER

            paragraphs = content.split("\n\n")
            for p in paragraphs:
                cleaned = p.strip()
                if cleaned:
                    if cleaned.startswith("# "):
                        subhead = doc.add_heading(cleaned[2:], level=2)
                        subhead.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    elif cleaned.startswith("## "):
                        subhead = doc.add_heading(cleaned[3:], level=3)
                    else:
                        p_para = doc.add_paragraph(cleaned)
                        p_para.paragraph_format.first_line_indent = Inches(0.3)

        # Back Matter (if present)
        if front_matter and ("about_author" in front_matter or "acknowledgments" in front_matter):
            if "acknowledgments" in front_matter:
                doc.add_page_break()
                ack_head = doc.add_heading("Acknowledgments", level=2)
                ack_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for block in front_matter["acknowledgments"].split("\n\n"):
                    if block.strip():
                        doc.add_paragraph(block.strip())

            if "about_author" in front_matter:
                doc.add_page_break()
                bio_head = doc.add_heading("About the Author", level=2)
                bio_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for block in front_matter["about_author"].split("\n\n"):
                    if block.strip():
                        doc.add_paragraph(block.strip())

        doc.save(output_path)
        return output_path
